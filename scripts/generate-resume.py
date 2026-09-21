#!/usr/bin/env python3
"""Regenerate public resume.pdf and the downloadable DOCX from resume_content.py."""

from __future__ import annotations

import sys
from pathlib import Path

from fpdf import FPDF
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))

import resume_content as C  # noqa: E402

PDF_PATH = ROOT / "public" / "resume.pdf"
DOCX_PATH = ROOT / "public" / "assets" / "Sunkara-Dineshgopi-Resume.docx"
TXT_PATH = ROOT / "public" / "assets" / "Sunkara-Dineshgopi-Resume.txt"

NAVY = (20, 32, 56)
RULE = (40, 56, 82)
MUTED = (70, 78, 90)
BODY = (28, 32, 38)


def ascii_dash(text: str) -> str:
    return (
        text.replace("–", "-")
        .replace("—", "-")
        .replace("’", "'")
        .replace("“", '"')
        .replace("”", '"')
    )


class ResumePDF(FPDF):
    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(*MUTED)
        self.cell(0, 8, ascii_dash(C.FOOTER), align="C")


def add_rule(pdf: ResumePDF, gap_before=0.8, gap_after=1.5):
    pdf.ln(gap_before)
    y = pdf.get_y()
    pdf.set_draw_color(*RULE)
    pdf.set_line_width(0.3)
    pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
    pdf.ln(gap_after)


def section_label(pdf: ResumePDF, title: str):
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_text_color(*NAVY)
    pdf.cell(0, 5.2, title, new_x="LMARGIN", new_y="NEXT")
    add_rule(pdf, gap_before=0.2, gap_after=1.4)


def write_pdf() -> None:
    pdf = ResumePDF(format="Letter", unit="mm")
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.set_margins(14.5, 11.5, 14.5)
    pdf.add_page()
    pdf.set_title("Sunkara Dineshgopi - Latest resume")
    pdf.set_author("Sunkara Dineshgopi")

    pdf.set_font("Helvetica", "B", 18)
    pdf.set_text_color(*NAVY)
    pdf.cell(0, 7, C.NAME, new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(*BODY)
    pdf.cell(0, 4.4, ascii_dash(f"{C.ROLE}  |  Latest resume  |  currently working"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(*MUTED)
    pdf.cell(0, 4.1, ascii_dash(f"{C.LOCATION}  |  {C.PHONE}  |  {C.EMAIL}"), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 4.1, f"{C.LINKEDIN_SHORT}  |  {C.GITHUB_SHORT}", new_x="LMARGIN", new_y="NEXT")

    section_label(pdf, "SUMMARY")
    pdf.set_font("Helvetica", "", 8.7)
    pdf.set_text_color(*BODY)
    pdf.multi_cell(0, 3.85, ascii_dash(C.SUMMARY))

    pdf.ln(1.4)
    section_label(pdf, "EXPERIENCE")
    for job in C.EXPERIENCE:
        pdf.set_font("Helvetica", "B", 9.4)
        pdf.set_text_color(*NAVY)
        company = ascii_dash(f"{job['company']} | {job['location']}")
        dates = ascii_dash(job["dates"])
        page_width = pdf.w - pdf.l_margin - pdf.r_margin
        pdf.cell(page_width - 40, 4.4, company)
        pdf.set_font("Helvetica", "B", 8.6)
        pdf.cell(40, 4.4, dates, align="R", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "I", 8.7)
        pdf.set_text_color(*BODY)
        pdf.cell(0, 4.0, ascii_dash(job["role"]), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 8.5)
        pdf.set_text_color(*BODY)
        for bullet in job["bullets"]:
            pdf.multi_cell(0, 3.7, ascii_dash(f"-  {bullet}"))
            pdf.ln(0.15)
        pdf.ln(0.85)

    pdf.add_page()
    section_label(pdf, "EDUCATION")
    for item in C.EDUCATION:
        pdf.set_font("Helvetica", "B", 9.4)
        pdf.set_text_color(*NAVY)
        page_width = pdf.w - pdf.l_margin - pdf.r_margin
        pdf.cell(page_width - 40, 4.4, ascii_dash(item["credential"]))
        pdf.set_font("Helvetica", "B", 8.6)
        pdf.cell(40, 4.4, ascii_dash(item["dates"]), align="R", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 8.7)
        pdf.set_text_color(*BODY)
        pdf.cell(0, 4.0, ascii_dash(item["school"]), new_x="LMARGIN", new_y="NEXT")

    pdf.ln(1.2)
    section_label(pdf, "SKILLS")
    for label, detail in C.SKILLS:
        pdf.set_font("Helvetica", "B", 8.5)
        pdf.set_text_color(*NAVY)
        pdf.write(3.7, ascii_dash(f"{label}: "))
        pdf.set_font("Helvetica", "", 8.3)
        pdf.set_text_color(*BODY)
        pdf.multi_cell(0, 3.85, ascii_dash(detail))
        pdf.ln(0.55)

    PDF_PATH.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(PDF_PATH))


def set_run_font(run, name="Calibri", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "283852")
    pBdr.append(bottom)
    pPr.append(pBdr)


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True, color=(20, 32, 56))
    add_bottom_border(p)
    return p


def write_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.line_spacing = 1.08

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(2)
    run = name.add_run(C.NAME)
    set_run_font(run, size=20, bold=True, color=(20, 32, 56))

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(1)
    run = contact.add_run(f"M: {C.PHONE}   |   E: {C.EMAIL}")
    set_run_font(run, size=10, color=(70, 78, 90))

    links = doc.add_paragraph()
    links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links.paragraph_format.space_after = Pt(6)
    run = links.add_run(f"{C.LINKEDIN_SHORT}   |   {C.GITHUB_SHORT}")
    set_run_font(run, size=10, color=(70, 78, 90))

    heading(doc, "Summary")
    summary = doc.add_paragraph()
    summary.paragraph_format.space_after = Pt(6)
    run = summary.add_run(C.SUMMARY)
    set_run_font(run, size=10.5, color=(28, 32, 38))

    heading(doc, "Skills")
    for label, detail in C.SKILLS:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.25)
        bold = p.add_run(f"{label}: ")
        set_run_font(bold, size=10.5, bold=True, color=(20, 32, 56))
        rest = p.add_run(detail)
        set_run_font(rest, size=10.5, color=(28, 32, 38))

    heading(doc, "Experience")
    for job in C.EXPERIENCE:
        header = doc.add_paragraph()
        header.paragraph_format.space_before = Pt(8)
        header.paragraph_format.space_after = Pt(0)
        tab_stops = header.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Inches(7.1), WD_TAB_ALIGNMENT.RIGHT)
        left = header.add_run(f"{job['company']} | {job['location']}")
        set_run_font(left, size=11, bold=True, color=(20, 32, 56))
        right = header.add_run(f"\t{job['dates']}")
        set_run_font(right, size=10.5, bold=True, color=(20, 32, 56))

        role = doc.add_paragraph()
        role.paragraph_format.space_after = Pt(2)
        run = role.add_run(job["role"])
        set_run_font(run, size=10.5, color=(28, 32, 38))

        for bullet in job["bullets"]:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.25)
            run = p.add_run(bullet)
            set_run_font(run, size=10, color=(28, 32, 38))

    heading(doc, "Education")
    for item in C.EDUCATION:
        header = doc.add_paragraph()
        header.paragraph_format.space_before = Pt(6)
        header.paragraph_format.space_after = Pt(0)
        tab_stops = header.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Inches(7.1), WD_TAB_ALIGNMENT.RIGHT)
        left = header.add_run(item["credential"])
        set_run_font(left, size=11, bold=True, color=(20, 32, 56))
        right = header.add_run(f"\t{item['dates']}")
        set_run_font(right, size=10, bold=True, color=(20, 32, 56))
        school = doc.add_paragraph()
        run = school.add_run(item["school"])
        set_run_font(run, size=10.5, color=(28, 32, 38))

    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(DOCX_PATH))


def write_txt() -> None:
    TXT_PATH.parent.mkdir(parents=True, exist_ok=True)
    TXT_PATH.write_text(C.plain_text() + "\n", encoding="utf-8")


if __name__ == "__main__":
    write_pdf()
    write_docx()
    write_txt()
    print(f"Wrote {PDF_PATH.relative_to(ROOT)}")
    print(f"Wrote {DOCX_PATH.relative_to(ROOT)}")
    print(f"Wrote {TXT_PATH.relative_to(ROOT)}")
